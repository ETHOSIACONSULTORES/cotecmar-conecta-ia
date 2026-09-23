from io import BytesIO
from pathlib import Path
from pypdf import PdfReader
from docx import Document

MAX_TEXT_CHARS = 45000

class DocumentError(Exception):
    pass

def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            parts = []
            for page in reader.pages:
                txt = page.extract_text() or ""
                if txt.strip():
                    parts.append(txt)
                if sum(len(x) for x in parts) >= MAX_TEXT_CHARS:
                    break
            text = "\n".join(parts)
        elif suffix == ".docx":
            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            # include tables
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text.strip() for cell in row.cells)
        else:
            raise DocumentError("Formato no soportado. Use PDF o DOCX.")
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"No fue posible leer el documento: {exc}") from exc
    return text[:MAX_TEXT_CHARS].strip()
